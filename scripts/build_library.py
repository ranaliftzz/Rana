#!/usr/bin/env python3
"""
Rana — exercise library repair + JSON export (PRD rana-prd-v1-910, section 4).

What this fixes: `workout-data/workout_splits_database.xlsx` has an `exercises_master`
sheet where every training-day row names a primary exercise plus two alternatives as
free text. 51 of the 85 distinct alternative names didn't exist in `exercise_library`,
so they'd fail silently the moment the app tried to show a swap option. This script:

  1. Takes the 56 exercises already in `exercise_library` as-is (already keyed).
  2. Adds NEW_ENTRIES below: full library rows for the ~40 alternative names that are
     genuinely distinct exercises, each with a hand-assigned exercise_key.
  3. Applies ALIASES below: alternative names that are just naming drift for an
     exercise the library already has (e.g. "Hack Squat" vs "Machine Hack Squat"),
     mapped onto the existing key instead of getting a new one.
  4. Resolves every exercise_name / alternative_1 / alternative_2 string anywhere in
     exercises_master to a real exercise_key, and fails loudly (raises) if anything
     doesn't resolve — same check `validate_library.py` re-runs standalone.
  5. Converts exercises_master's flat rows into the PRD's slot / slot_option model:
     each row already *is* one slot (a movement-pattern requirement) with up to three
     ranked options (primary, alternative_1, alternative_2), so the conversion is
     mechanical, not a redesign.
  6. Exports data/exercises.json, data/splits.json, data/workouts.json, data/slots.json.

Run: python3 scripts/build_library.py
Requires: openpyxl (only for reading the .xlsx — the generated JSON has no dependency).
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "workout-data" / "workout_splits_database.xlsx"
DATA_DIR = ROOT / "data"

# ---------------------------------------------------------------------------
# ALIASES — alternative names that are the SAME exercise as an existing library
# entry, just written differently. Confirmed by hand (2026-09-11), not fuzzy-matched:
# 5 named directly in the PRD's decision log, plus 3 more found by cross-checking
# equipment + movement_pattern against the library and confirmed with the user.
# ---------------------------------------------------------------------------
ALIASES = {
    "Hack Squat": "hack_squat",
    "Lat Pulldown": "lat_pulldown",
    "Seated Cable Row": "seated_cable_row",
    "Incline Barbell Bench Press": "incline_bb",
    "Dumbbell Bench Press": "flat_db_press",
    "Cable Fly": "cable_fly",
    "Walking Lunge": "walking_lunge",
    "Cable Lat Prayer": "straight_arm_pd",
}

# ---------------------------------------------------------------------------
# NEW_ENTRIES — the ~40 alternative names left over after ALIASES, each a genuinely
# distinct exercise (different equipment, different movement, or both) with no
# existing library entry. Full schema per PRD section 5.1. Equipment-variant policy
# confirmed with the user 2026-09-11: distinct equipment/attachment always gets its
# own key (relevant to the swap feature), never merged into the closest match.
#
# Drafted with standard exercise-science classifications, matching the voice and
# controlled vocabulary of the existing 56 entries. Flag anything that reads wrong
# to the fitness-expert eye — these haven't been programming-reviewed yet.
# ---------------------------------------------------------------------------
NEW_ENTRIES = [
    dict(exercise_key="assisted_pullup", exercise_name="Assisted Pull-Up", muscle_group="Back",
         movement_pattern="Vertical Pull", target_muscle_primary="Latissimus Dorsi",
         target_muscle_secondary="Biceps Brachii; Teres Major", equipment="Machine",
         is_compound=True, is_unilateral=False,
         coaching_cue="Set the counterweight to the lightest assistance that still lets you complete every rep with control — the goal is progressing toward less assistance, not resting on the pad."),
    dict(exercise_key="band_pull_apart", exercise_name="Band Pull-Apart", muscle_group="Back",
         movement_pattern="Horizontal Abduction", target_muscle_primary="Posterior Deltoid",
         target_muscle_secondary="Rhomboids; Mid-Trapezius", equipment="Band",
         is_compound=False, is_unilateral=False,
         coaching_cue="Arms straight, pull the band to the chest by squeezing the shoulder blades together, not by bending the elbows."),
    dict(exercise_key="bayesian_cable_curl", exercise_name="Bayesian Cable Curl", muscle_group="Biceps",
         movement_pattern="Elbow Flexion (Stretched)", target_muscle_primary="Biceps Brachii",
         target_muscle_secondary="Brachialis", equipment="Cable",
         is_compound=False, is_unilateral=True,
         coaching_cue="Let the arm trail behind the torso at the bottom for a deep stretch, then curl without swinging the shoulder forward."),
    dict(exercise_key="bench_dip", exercise_name="Bench Dip", muscle_group="Triceps",
         movement_pattern="Elbow Extension", target_muscle_primary="Triceps Brachii",
         target_muscle_secondary="Anterior Deltoid; Sternal Pectoralis", equipment="Bodyweight",
         is_compound=False, is_unilateral=False,
         coaching_cue="Keep hips close to the bench and elbows pointed back, not flared out, to keep the load on the triceps."),
    dict(exercise_key="cable_crunch", exercise_name="Cable Crunch", muscle_group="Core",
         movement_pattern="Trunk Flexion", target_muscle_primary="Rectus Abdominis",
         target_muscle_secondary="Obliques", equipment="Cable",
         is_compound=False, is_unilateral=False,
         coaching_cue="Round the spine and crunch toward the hips using the abs, not the hip flexors — the hips barely move."),
    dict(exercise_key="cable_pull_through", exercise_name="Cable Pull-Through", muscle_group="Glutes",
         movement_pattern="Hip Hinge", target_muscle_primary="Gluteus Maximus",
         target_muscle_secondary="Hamstrings", equipment="Cable",
         is_compound=True, is_unilateral=False,
         coaching_cue="Hinge at the hips with a soft knee bend, push the hips back until you feel the stretch, then drive them forward to finish."),
    dict(exercise_key="cable_reverse_fly", exercise_name="Cable Reverse Fly", muscle_group="Shoulders",
         movement_pattern="Horizontal Abduction", target_muscle_primary="Posterior Deltoid",
         target_muscle_secondary="Rhomboids; Mid-Trapezius", equipment="Cable",
         is_compound=False, is_unilateral=False,
         coaching_cue="Cross the cables in front of you and sweep the arms out and back with a slight elbow bend, leading with the elbows."),
    dict(exercise_key="cable_shrug", exercise_name="Cable Shrug", muscle_group="Traps",
         movement_pattern="Scapular Elevation", target_muscle_primary="Upper Trapezius",
         target_muscle_secondary="Levator Scapulae", equipment="Cable",
         is_compound=False, is_unilateral=False,
         coaching_cue="Shrug straight up toward the ears, pause, and lower under control — no rolling the shoulders."),
    dict(exercise_key="captains_chair_raise", exercise_name="Captain's Chair Knee Raise", muscle_group="Core",
         movement_pattern="Trunk Flexion", target_muscle_primary="Rectus Abdominis",
         target_muscle_secondary="Hip Flexors", equipment="Machine",
         is_compound=False, is_unilateral=False,
         coaching_cue="Press the back into the pad and curl the hips up toward the ribs, not just lift the knees."),
    dict(exercise_key="dead_bug", exercise_name="Dead Bug", muscle_group="Core",
         movement_pattern="Anti-Extension", target_muscle_primary="Rectus Abdominis",
         target_muscle_secondary="Obliques; Hip Flexors", equipment="Bodyweight",
         is_compound=False, is_unilateral=True,
         coaching_cue="Press the low back into the floor and keep it pinned there through the entire rep as the opposite arm and leg extend."),
    dict(exercise_key="decline_db_fly", exercise_name="Decline Dumbbell Fly", muscle_group="Chest",
         movement_pattern="Horizontal Adduction", target_muscle_primary="Lower Pectoralis",
         target_muscle_secondary="Anterior Deltoid", equipment="Dumbbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="Slight bend in the elbows held constant, arc the dumbbells together as if hugging a barrel, feeling it in the lower chest."),
    dict(exercise_key="decline_db_press", exercise_name="Decline Dumbbell Press", muscle_group="Chest",
         movement_pattern="Decline Push", target_muscle_primary="Lower Pectoralis",
         target_muscle_secondary="Triceps Brachii; Anterior Deltoid", equipment="Dumbbell",
         is_compound=True, is_unilateral=False,
         coaching_cue="Press up and slightly back toward the hips, tracking the natural decline angle rather than straight up."),
    dict(exercise_key="db_curl", exercise_name="Dumbbell Curl", muscle_group="Biceps",
         movement_pattern="Elbow Flexion", target_muscle_primary="Biceps Brachii",
         target_muscle_secondary="Brachialis", equipment="Dumbbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="Keep elbows pinned to your sides and rotate to a full supinated grip at the top without swinging."),
    dict(exercise_key="db_fly", exercise_name="Dumbbell Fly", muscle_group="Chest",
         movement_pattern="Horizontal Adduction", target_muscle_primary="Sternal Pectoralis",
         target_muscle_secondary="Anterior Deltoid", equipment="Dumbbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="Soft bend in the elbows held constant, lower until you feel a stretch across the chest, then arc back up."),
    dict(exercise_key="db_oh_extension", exercise_name="Dumbbell Overhead Extension", muscle_group="Triceps",
         movement_pattern="Elbow Extension (Overhead)", target_muscle_primary="Triceps Brachii (long head)",
         target_muscle_secondary="", equipment="Dumbbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="Keep the elbows tucked in close to the head and only the forearms move — the upper arms stay still."),
    dict(exercise_key="db_preacher_curl", exercise_name="Dumbbell Preacher Curl", muscle_group="Biceps",
         movement_pattern="Elbow Flexion (Shortened)", target_muscle_primary="Biceps Brachii",
         target_muscle_secondary="Brachialis", equipment="Dumbbell",
         is_compound=False, is_unilateral=True,
         coaching_cue="Let the arm fully straighten at the bottom against the pad, and stop just short of lockout at the top to keep tension on the biceps."),
    dict(exercise_key="db_rear_delt_fly", exercise_name="Dumbbell Rear Delt Fly", muscle_group="Shoulders",
         movement_pattern="Horizontal Abduction", target_muscle_primary="Posterior Deltoid",
         target_muscle_secondary="Rhomboids", equipment="Dumbbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="Hinge forward with a flat back, and raise the dumbbells out to the sides leading with the pinkies, not the thumbs."),
    dict(exercise_key="db_rdl", exercise_name="Dumbbell Romanian Deadlift", muscle_group="Hamstrings",
         movement_pattern="Hip Hinge", target_muscle_primary="Hamstrings",
         target_muscle_secondary="Gluteus Maximus; Erector Spinae", equipment="Dumbbell",
         is_compound=True, is_unilateral=False,
         coaching_cue="Push the hips straight back with a soft knee bend, keeping the dumbbells close to the legs, until you feel the hamstring stretch."),
    dict(exercise_key="db_shrug", exercise_name="Dumbbell Shrug", muscle_group="Traps",
         movement_pattern="Scapular Elevation", target_muscle_primary="Upper Trapezius",
         target_muscle_secondary="Levator Scapulae", equipment="Dumbbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="Shrug straight up, pause at the top, and lower under control — don't roll the shoulders forward or back."),
    dict(exercise_key="ez_bar_curl", exercise_name="EZ-Bar Curl", muscle_group="Biceps",
         movement_pattern="Elbow Flexion", target_muscle_primary="Biceps Brachii",
         target_muscle_secondary="Brachialis", equipment="Barbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="The cambered grip takes strain off the wrists — keep elbows still at your sides through the full range."),
    dict(exercise_key="glute_bridge", exercise_name="Glute Bridge", muscle_group="Glutes",
         movement_pattern="Hip Extension", target_muscle_primary="Gluteus Maximus",
         target_muscle_secondary="Hamstrings", equipment="Bodyweight",
         is_compound=True, is_unilateral=False,
         coaching_cue="Drive through the heels and squeeze the glutes hard at the top rather than arching the lower back to finish the rep."),
    dict(exercise_key="goblet_squat", exercise_name="Goblet Squat", muscle_group="Quads",
         movement_pattern="Squat (Knee Dominant)", target_muscle_primary="Quadriceps",
         target_muscle_secondary="Gluteus Maximus; Adductors", equipment="Dumbbell",
         is_compound=True, is_unilateral=False,
         coaching_cue="Hold the dumbbell vertically at chest height and let it counterbalance you into a deep, upright-torso squat."),
    dict(exercise_key="incline_machine_press", exercise_name="Incline Machine Press", muscle_group="Chest",
         movement_pattern="Incline Push", target_muscle_primary="Upper Pectoralis",
         target_muscle_secondary="Anterior Deltoid; Triceps Brachii", equipment="Machine",
         is_compound=True, is_unilateral=False,
         coaching_cue="Set the seat so the handles line up with the upper chest, and press up and slightly back along the fixed path."),
    dict(exercise_key="incline_smith_press", exercise_name="Incline Smith Machine Press", muscle_group="Chest",
         movement_pattern="Incline Push", target_muscle_primary="Upper Pectoralis",
         target_muscle_secondary="Anterior Deltoid; Triceps Brachii", equipment="Machine",
         is_compound=True, is_unilateral=False,
         coaching_cue="The fixed bar path removes the need to stabilize side-to-side — put that saved effort into driving through the upper chest."),
    dict(exercise_key="leg_press_calf_raise", exercise_name="Leg Press Calf Raise", muscle_group="Calves",
         movement_pattern="Ankle Plantarflexion", target_muscle_primary="Gastrocnemius",
         target_muscle_secondary="Soleus", equipment="Machine",
         is_compound=False, is_unilateral=False,
         coaching_cue="Legs stay nearly straight on the sled — press through the balls of the feet and let the heels drop for a full stretch each rep."),
    dict(exercise_key="leg_press_calf_raise_bent_knee", exercise_name="Leg Press Calf Raise (Bent Knee)", muscle_group="Calves",
         movement_pattern="Ankle Plantarflexion", target_muscle_primary="Soleus",
         target_muscle_secondary="Gastrocnemius", equipment="Machine",
         is_compound=False, is_unilateral=False,
         coaching_cue="Keep the knees bent throughout — this takes the gastrocnemius out and shifts the load onto the soleus underneath."),
    dict(exercise_key="machine_assisted_dip", exercise_name="Machine Assisted Dip", muscle_group="Chest",
         movement_pattern="Decline Push", target_muscle_primary="Lower Pectoralis",
         target_muscle_secondary="Triceps Brachii; Anterior Deltoid", equipment="Machine",
         is_compound=True, is_unilateral=False,
         coaching_cue="Use the least assistance that still lets you complete every rep with a full stretch at the bottom — assistance goes down over time, not up."),
    dict(exercise_key="machine_hip_thrust", exercise_name="Machine Hip Thrust", muscle_group="Glutes",
         movement_pattern="Hip Extension", target_muscle_primary="Gluteus Maximus",
         target_muscle_secondary="Hamstrings", equipment="Machine",
         is_compound=True, is_unilateral=False,
         coaching_cue="Drive the pad up by squeezing the glutes, not by arching through the lower back."),
    dict(exercise_key="machine_preacher_curl", exercise_name="Machine Preacher Curl", muscle_group="Biceps",
         movement_pattern="Elbow Flexion (Shortened)", target_muscle_primary="Biceps Brachii",
         target_muscle_secondary="Brachialis", equipment="Machine",
         is_compound=False, is_unilateral=False,
         coaching_cue="Let the arms fully extend at the bottom against the pad and stop just short of lockout at the top."),
    dict(exercise_key="machine_pullover", exercise_name="Machine Pullover", muscle_group="Back",
         movement_pattern="Shoulder Extension", target_muscle_primary="Latissimus Dorsi",
         target_muscle_secondary="Sternal Pectoralis; Long Head Triceps", equipment="Machine",
         is_compound=False, is_unilateral=False,
         coaching_cue="Lead with the elbows on the way down and think about pulling the bar with the lats, not the arms."),
    dict(exercise_key="machine_shoulder_press", exercise_name="Machine Shoulder Press", muscle_group="Shoulders",
         movement_pattern="Vertical Push", target_muscle_primary="Anterior Deltoid",
         target_muscle_secondary="Triceps Brachii; Medial Deltoid", equipment="Machine",
         is_compound=True, is_unilateral=False,
         coaching_cue="Set the seat so the handles start level with the shoulders, and press straight up along the fixed path without shrugging."),
    dict(exercise_key="nordic_curl", exercise_name="Nordic Hamstring Curl", muscle_group="Hamstrings",
         movement_pattern="Knee Flexion", target_muscle_primary="Hamstrings",
         target_muscle_secondary="Gluteus Maximus", equipment="Bodyweight",
         is_compound=False, is_unilateral=False,
         coaching_cue="Anchor the ankles and lower the torso as slowly as possible, resisting with the hamstrings the whole way down."),
    dict(exercise_key="pushup", exercise_name="Push-Up", muscle_group="Chest",
         movement_pattern="Horizontal Push", target_muscle_primary="Sternal Pectoralis",
         target_muscle_secondary="Triceps Brachii; Anterior Deltoid", equipment="Bodyweight",
         is_compound=True, is_unilateral=False,
         coaching_cue="Keep a straight line from shoulders to ankles and lower until the chest is just above the floor."),
    dict(exercise_key="reverse_crunch", exercise_name="Reverse Crunch", muscle_group="Core",
         movement_pattern="Trunk Flexion", target_muscle_primary="Rectus Abdominis (lower)",
         target_muscle_secondary="Hip Flexors", equipment="Bodyweight",
         is_compound=False, is_unilateral=False,
         coaching_cue="Curl the hips up off the floor toward the ribs using the lower abs, not momentum from swinging the legs."),
    dict(exercise_key="reverse_curl", exercise_name="Reverse Curl", muscle_group="Biceps",
         movement_pattern="Elbow Flexion", target_muscle_primary="Brachioradialis",
         target_muscle_secondary="Biceps Brachii", equipment="Barbell",
         is_compound=False, is_unilateral=False,
         coaching_cue="Pronated (overhand) grip throughout — keep the wrists locked flat and let the forearms and brachialis do the work."),
    dict(exercise_key="reverse_lunge", exercise_name="Reverse Lunge", muscle_group="Quads",
         movement_pattern="Lunge (Unilateral)", target_muscle_primary="Quadriceps",
         target_muscle_secondary="Gluteus Maximus", equipment="Dumbbell",
         is_compound=True, is_unilateral=True,
         coaching_cue="Step backward into the lunge and push through the front heel to return to standing — easier on the front knee than stepping forward."),
    dict(exercise_key="reverse_nordic_curl", exercise_name="Reverse Nordic Curl", muscle_group="Quads",
         movement_pattern="Knee Extension", target_muscle_primary="Quadriceps",
         target_muscle_secondary="Hip Flexors", equipment="Bodyweight",
         is_compound=False, is_unilateral=False,
         coaching_cue="Kneeling with hips locked long, lean straight back from the knees as far as control allows, then pull back up."),
    dict(exercise_key="sa_db_row", exercise_name="Single-Arm Dumbbell Row", muscle_group="Back",
         movement_pattern="Horizontal Pull", target_muscle_primary="Latissimus Dorsi",
         target_muscle_secondary="Rhomboids; Biceps Brachii", equipment="Dumbbell",
         is_compound=True, is_unilateral=True,
         coaching_cue="Brace the free hand on a bench, keep the torso still, and pull the elbow up and back rather than rotating the trunk to finish."),
    dict(exercise_key="sissy_squat", exercise_name="Sissy Squat", muscle_group="Quads",
         movement_pattern="Knee Extension", target_muscle_primary="Quadriceps",
         target_muscle_secondary="", equipment="Bodyweight",
         is_compound=False, is_unilateral=False,
         coaching_cue="Rise onto the toes and lean straight back from the knees, keeping hips extended, so the quads take the full load."),
    dict(exercise_key="smith_calf_raise", exercise_name="Smith Machine Calf Raise", muscle_group="Calves",
         movement_pattern="Ankle Plantarflexion", target_muscle_primary="Gastrocnemius",
         target_muscle_secondary="Soleus", equipment="Machine",
         is_compound=False, is_unilateral=False,
         coaching_cue="Stand with the balls of the feet on a plate for a deep stretch, and press straight up through the toes."),
    dict(exercise_key="smith_seated_calf_raise", exercise_name="Smith Machine Seated Calf Raise", muscle_group="Calves",
         movement_pattern="Ankle Plantarflexion", target_muscle_primary="Soleus",
         target_muscle_secondary="Gastrocnemius", equipment="Machine",
         is_compound=False, is_unilateral=False,
         coaching_cue="Knees bent at 90 degrees under the bar — this isolates the soleus underneath the gastrocnemius."),
    dict(exercise_key="step_up", exercise_name="Step-Up", muscle_group="Quads",
         movement_pattern="Lunge (Unilateral)", target_muscle_primary="Quadriceps",
         target_muscle_secondary="Gluteus Maximus", equipment="Dumbbell",
         is_compound=True, is_unilateral=True,
         coaching_cue="Drive through the whole foot on the box, not a push off the trailing leg, to keep the working leg doing the work."),
    dict(exercise_key="straight_bar_pushdown", exercise_name="Straight-Bar Pushdown", muscle_group="Triceps",
         movement_pattern="Elbow Extension", target_muscle_primary="Triceps Brachii",
         target_muscle_secondary="", equipment="Cable",
         is_compound=False, is_unilateral=False,
         coaching_cue="Keep the elbows pinned to your sides and only the forearms move — don't let the elbows drift forward as you push down."),
]


def load_existing_library(wb):
    ws = wb["exercise_library"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    entries = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        entries.append(dict(
            exercise_key=row[idx["exercise_key"]],
            exercise_name=row[idx["exercise_name"]],
            muscle_group=row[idx["muscle_group"]],
            movement_pattern=row[idx["movement_pattern"]],
            target_muscle_primary=row[idx["target_muscle_primary"]],
            target_muscle_secondary=row[idx["target_muscle_secondary"]] or "",
            equipment=row[idx["equipment"]],
            is_compound=str(row[idx["is_compound"]]).upper() == "TRUE",
            is_unilateral=str(row[idx["is_unilateral"]]).upper() == "TRUE",
            coaching_cue=row[idx["coaching_cues"]],
        ))
    return entries


def build_exercise_library(wb):
    entries = load_existing_library(wb) + NEW_ENTRIES
    keys_seen = set()
    for e in entries:
        if e["exercise_key"] in keys_seen:
            raise ValueError(f"duplicate exercise_key: {e['exercise_key']}")
        keys_seen.add(e["exercise_key"])

    name_to_key = {e["exercise_name"]: e["exercise_key"] for e in entries}
    for alias_name, target_key in ALIASES.items():
        if target_key not in keys_seen:
            raise ValueError(f"alias '{alias_name}' points at unknown key '{target_key}'")
        name_to_key[alias_name] = target_key

    return entries, name_to_key


def build_splits(wb):
    ws = wb["splits_overview"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    splits = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["split_id"]] is None:
            continue
        splits.append({h: row[idx[h]] for h in header})
    return splits


def build_workouts(wb):
    ws = wb["workout_days"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    workouts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["workout_id"]] is None:
            continue
        workouts.append({h: row[idx[h]] for h in header})
    return workouts


def build_slots(wb, name_to_key):
    """Each exercises_master row IS a slot: a movement-pattern requirement with up
    to 3 ranked options (primary, alternative_1, alternative_2). slot_id reuses the
    sheet's own exercise_instance_id since it's already a stable unique int."""
    ws = wb["exercises_master"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}

    slots = []
    slot_options = []
    unresolved = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["exercise_instance_id"]] is None:
            continue
        slot_id = row[idx["exercise_instance_id"]]

        def resolve(name):
            if not name:
                return None
            key = name_to_key.get(name.strip())
            if key is None:
                unresolved.append((slot_id, name))
            return key

        primary_key = resolve(row[idx["exercise_name"]])
        alt1_key = resolve(row[idx["alternative_1"]])
        alt2_key = resolve(row[idx["alternative_2"]])

        slots.append(dict(
            slot_id=slot_id,
            workout_id=row[idx["workout_id"]],
            slot_order=row[idx["exercise_order"]],
            movement_pattern=row[idx["movement_pattern"]],
            working_sets=row[idx["working_sets"]],
            rep_range=row[idx["rep_range"]],
            target_rpe=row[idx["target_rpe"]],
            rest_seconds=row[idx["rest_seconds"]],
            tempo=row[idx["tempo"]],
            coaching_cue=row[idx["coaching_cues"]],
        ))
        if primary_key:
            slot_options.append(dict(slot_id=slot_id, exercise_key=primary_key, option_rank=1))
        if alt1_key:
            slot_options.append(dict(slot_id=slot_id, exercise_key=alt1_key, option_rank=2))
        if alt2_key:
            slot_options.append(dict(slot_id=slot_id, exercise_key=alt2_key, option_rank=3))

    if unresolved:
        lines = "\n".join(f"  slot {sid}: '{name}'" for sid, name in unresolved)
        raise ValueError(f"{len(unresolved)} exercise reference(s) did not resolve to a library key:\n{lines}")

    return slots, slot_options


def main():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    library, name_to_key = build_exercise_library(wb)
    splits = build_splits(wb)
    workouts = build_workouts(wb)
    slots, slot_options = build_slots(wb, name_to_key)

    DATA_DIR.mkdir(exist_ok=True)
    (DATA_DIR / "exercises.json").write_text(json.dumps(library, indent=2) + "\n")
    (DATA_DIR / "splits.json").write_text(json.dumps(splits, indent=2) + "\n")
    (DATA_DIR / "workouts.json").write_text(json.dumps(workouts, indent=2) + "\n")
    (DATA_DIR / "slots.json").write_text(json.dumps({"slots": slots, "slot_options": slot_options}, indent=2) + "\n")

    print(f"exercises.json: {len(library)} entries ({len(library) - len(NEW_ENTRIES)} existing + {len(NEW_ENTRIES)} new)")
    print(f"splits.json: {len(splits)} splits")
    print(f"workouts.json: {len(workouts)} workout days")
    print(f"slots.json: {len(slots)} slots, {len(slot_options)} slot_options")
    print(f"aliases resolved: {len(ALIASES)}")
    print("All exercise references resolved cleanly.")


if __name__ == "__main__":
    try:
        main()
    except ValueError as e:
        print(f"BUILD FAILED: {e}", file=sys.stderr)
        sys.exit(1)
